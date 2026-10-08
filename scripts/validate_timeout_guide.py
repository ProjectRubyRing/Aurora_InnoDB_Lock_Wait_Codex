"""Check workbook integrity, independent formula evaluation and measured layout.

This is artifact validation, not a JBoss/Aurora integration test. Preview PNGs
are drawn from workbook styles with Meiryo UI, not from the Excel COM engine.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from PIL import Image, ImageDraw
from pycel import ExcelCompiler

from build_timeout_guide import (BOOK_PATH, FONT_NAME, OUT, ROOT, WIDTHS,
                                 budget_results, effective_limit, measured_font,
                                 pixel_width, wrap_lines)
from guide_content import EFFECTIVE_CASES, INPUTS, SECTIONS, SOURCES


def close_equal(got, expected):
    if isinstance(expected, (int, float)):
        return isinstance(got, (int, float)) and math.isclose(got, expected, abs_tol=1e-8)
    return got == expected


def validate_formulas(wb, cached_wb):
    compiler = ExcelCompiler(filename=str(BOOK_PATH))
    name = next(n for n in wb.sheetnames if n.startswith("15_"))
    names = [f"{name}!E{r}" for r in range(6, 21)]
    results = []
    scenarios = [
        ("初期値", {}),
        ("個別Pなし・固定Q上書き", {3: 0}),
        ("SQL3本で予算不足", {11: 3}),
        ("JTA残り5秒・開始拒否予算", {1: 55}),
        ("JTA残り0・SQL実行拒否", {1: 60}),
        ("JTA残り負・SQL実行拒否", {1: 65}),
        ("JTAなし・固定Q", {3: 0, 5: "false"}),
        ("false・個別Pなし", {3: 0, 4: "false"}),
        ("false・個別P20", {3: 20, 4: "false"}),
        ("SQLもsocketも期限なし", {2: 0, 3: 0, 4: "false", 5: "false", 7: 0}),
        ("端数R2.3", {1: 57.7}),
        ("端数R0.2", {1: 59.8}),
        ("外側Hが短い", {15: 40}),
        ("socketがSQLより短い", {7: 1000}),
        ("XA非ゼロも実効なし", {16: 120}),
    ]
    base = [row[1] for row in INPUTS]
    # Generate and populate the entire dependency graph before set_value.
    compiler.evaluate(names)
    for label, changes in scenarios:
        values = list(base)
        for idx, val in changes.items():
            values[idx] = val
        for idx, val in enumerate(values):
            addr = f"{name}!B{idx + 6}"
            compiler.evaluate(addr)
            compiler.set_value(addr, val)
        expected = budget_results(values)
        actual = compiler.evaluate(names)
        mismatches = [(addr, got, want) for addr, got, want in zip(names, actual, expected)
                      if not close_equal(got, want)]
        if mismatches:
            raise AssertionError((label, mismatches))
        results.append({"scenario": label, "checked_results": len(names), "passed": True})
    init_expected = budget_results(base)
    for addr, want in zip(names, init_expected):
        cell = addr.split("!")[1]
        got = cached_wb[name][cell].value
        assert close_equal(got, want), ("cache", cell, got, want)
    for case in EFFECTIVE_CASES:
        sid, jta, switch, remaining, q, p, want, _ = case
        assert effective_limit(jta, switch, remaining, q, p) == want, sid
    return results


def validate_structure(wb, cached_wb, manifest):
    expected_names = sorted([s[0] for s in SECTIONS] +
                            ["00_はじめに", "09_実効値早見", "15_設定診断器", "21_出典と確認範囲", "22_動作イメージ"])
    assert wb.sheetnames == expected_names
    assert all(f.name == FONT_NAME for f in wb._fonts), [(f.name, f.sz) for f in wb._fonts]
    source_ids = {s[0] for s in SOURCES}
    for section in SECTIONS:
        for row in section[3]:
            assert set(re.findall(r"S\d{2}", row[-1])) <= source_ids, row
    counts = {"populated_cells": 0, "formulas": 0, "hyperlinks": 0, "layout_rows": 0}
    overflows = []
    for ws in wb:
        assert ws.freeze_panes is not None, ws.title
        assert ws.sheet_view.showGridLines is False, ws.title
        assert ws.page_setup.orientation == "landscape", ws.title
        # OOXML defaults fitToWidth to 1; XlsxWriter omits that attribute.
        assert ws.page_setup.fitToWidth in (None, 1), ws.title
        assert ws.page_setup.fitToHeight == 0, ws.title
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True, ws.title
        for row in ws:
            for cell in row:
                if cell.value is None:
                    continue
                counts["populated_cells"] += 1
                assert cell.font.name == FONT_NAME, (ws.title, cell.coordinate, cell.font.name)
                assert cell.alignment.wrap_text, (ws.title, cell.coordinate, "wrap")
                assert "\ufffd" not in str(cell.value), (ws.title, cell.coordinate, "encoding")
                if cell.data_type == "f":
                    counts["formulas"] += 1
                    assert cached_wb[ws.title][cell.coordinate].value is not None
                if cell.hyperlink:
                    counts["hyperlinks"] += 1
                    if cell.hyperlink.target:
                        assert cell.hyperlink.target.startswith("https://"), (ws.title, cell.coordinate)
                    elif cell.hyperlink.location:
                        target = cell.hyperlink.location.split("!")[0].strip("'")
                        assert target in wb.sheetnames, (ws.title, cell.coordinate, target)
        layout = manifest["layout"][ws.title]
        for row_index, settings in layout["rows"].items():
            index = int(row_index)
            counts["layout_rows"] += 1
            actual_height = ws.row_dimensions[index + 1].height
            assert actual_height and actual_height <= 409, (ws.title, index, actual_height)
            for entry in settings["cells"]:
                col = entry["column"]
                cell = ws.cell(index + 1, col + 1)
                value = cached_wb[ws.title].cell(index + 1, col + 1).value
                if value is None:
                    continue
                fontsize = cell.font.sz or 11
                lines = wrap_lines(str(value), pixel_width(layout["widths"][col]), fontsize, bool(cell.font.b))
                # Row height uses a conservative 17.5pt line pitch at 11pt.
                # Header lines at 10/11pt need the same room; avoid clipping.
                needed = len(lines) * 17.5 + 10
                if needed > actual_height + 0.1:
                    overflows.append({"sheet": ws.title, "row": index + 1,
                                      "cell": cell.coordinate, "height": actual_height,
                                      "estimated_needed": needed, "text": str(value)[:85]})
        for merge in layout["merged"]:
            width = sum(pixel_width(w) for w in layout["widths"][merge["first"]:merge["last"] + 1])
            cell = ws.cell(merge["row"] + 1, merge["first"] + 1)
            line_count = len(wrap_lines(str(cell.value or ""), width, cell.font.sz, bool(cell.font.b)))
            needed = line_count * (27 if cell.font.sz > 16 else 17.5) + 7
            if needed > merge["height"] + 0.1:
                overflows.append({"sheet": ws.title, "row": merge["row"] + 1,
                                  "height": merge["height"], "estimated_needed": needed,
                                  "text": str(cell.value)[:85]})
    assert not overflows, json.dumps(overflows, ensure_ascii=False, indent=2)
    assert counts["formulas"] == len(manifest["formulas"])
    return counts


def rgb(value, default):
    if value and value.type == "rgb" and isinstance(value.rgb, str):
        return "#" + value.rgb[-6:]
    return default


def preview(wb, cached_wb, manifest, name, max_row):
    ws = wb[name]
    data = cached_wb[name]
    layout = manifest["layout"][name]
    widths = [pixel_width(w) for w in layout["widths"]]
    x_edges = [20]
    for width in widths:
        x_edges.append(x_edges[-1] + width)
    y_edges = [20]
    for row in range(1, max_row + 1):
        y_edges.append(y_edges[-1] + round((ws.row_dimensions[row].height or 22) * 96 / 72))
    im = Image.new("RGB", (x_edges[-1] + 20, y_edges[-1] + 74), "#ECF0F4")
    draw = ImageDraw.Draw(im)
    merged_lookup = {}
    skip = set()
    for region in ws.merged_cells.ranges:
        merged_lookup[(region.min_row, region.min_col)] = (region.max_row, region.max_col)
        for row in range(region.min_row, region.max_row + 1):
            for col in range(region.min_col, region.max_col + 1):
                if (row, col) != (region.min_row, region.min_col):
                    skip.add((row, col))
    for row in range(1, max_row + 1):
        for col in range(1, len(widths) + 1):
            if (row, col) in skip:
                continue
            last_row, last_col = merged_lookup.get((row, col), (row, col))
            if last_row > max_row:
                continue
            cell = ws.cell(row, col)
            value = data.cell(row, col).value
            x0, x1 = x_edges[col - 1], x_edges[last_col]
            y0, y1 = y_edges[row - 1], y_edges[last_row]
            fill = rgb(cell.fill.fgColor, "#FFFFFF") if cell.fill.patternType else "#FFFFFF"
            draw.rectangle((x0, y0, x1, y1), fill=fill, outline="#D9E3EA")
            if value is None:
                continue
            fontsize = cell.font.sz or 11
            bold = bool(cell.font.b)
            font = measured_font(fontsize, bold)
            lines = wrap_lines(str(value), x1 - x0, fontsize, bold)
            line_pitch = round((27 if fontsize > 16 else 17.5) * 96 / 72)
            y_text = y0 + 8
            if cell.alignment.vertical == "center":
                y_text = y0 + max(6, (y1 - y0 - len(lines) * line_pitch) // 2)
            for line in lines:
                xpos = x0 + 9
                if cell.alignment.horizontal == "center":
                    xpos = x0 + max(7, (x1 - x0 - font.getlength(line)) / 2)
                draw.text((xpos, y_text), line, font=font, fill=rgb(cell.font.color, "#223C4D"))
                y_text += line_pitch
    caption = "書式・文字幅による確認用プレビュー（Excel表示エンジンのスクリーンショットではありません）"
    draw.text((20, y_edges[-1] + 20), caption, font=measured_font(10), fill="#536B7A")
    folder = OUT / "previews"
    folder.mkdir(exist_ok=True)
    path = folder / f"{name[:2]}_layout_preview.png"
    im.save(path)
    return str(path.relative_to(ROOT))


def main():
    manifest = json.loads((OUT / "workbook_validation_manifest.json").read_text(encoding="utf-8"))
    with ZipFile(BOOK_PATH) as z:
        assert z.testzip() is None
        names = z.namelist()
        assert not any("vbaProject" in n or "externalLink" in n for n in names)
    wb = openpyxl.load_workbook(BOOK_PATH, data_only=False)
    cached = openpyxl.load_workbook(BOOK_PATH, data_only=True)
    stats = validate_structure(wb, cached, manifest)
    scenarios = validate_formulas(wb, cached)
    previews = [preview(wb, cached, manifest, name, limit)
                for name, limit in [("00_はじめに", 10), ("01_結論と設定案", 9),
                                    ("15_設定診断器", 22), ("22_動作イメージ", 29)]]
    for path in (ROOT / "scripts").glob("*.py"):
        text = path.read_text(encoding="utf-8")
        assert "\ufffd" not in text
        assert all(line == line.rstrip() for line in text.splitlines()), path
    report = {"status": "passed", "workbook": BOOK_PATH.name,
              "sha256": hashlib.sha256(BOOK_PATH.read_bytes()).hexdigest(),
              "sheets": len(wb.sheetnames), "sources": len(SOURCES),
              "narrative_rows": sum(len(s[3]) for s in SECTIONS), **stats,
              "formula_scenarios": scenarios, "truth_table_cases": len(EFFECTIVE_CASES),
              "font": FONT_NAME, "measured_layout_checked": True,
              "cached_formula_results_checked": True, "preview_files": previews,
              "excel_com_display_check": "not performed: Excel COM activation failed with 0x80080005",
              "runtime_tests": "not performed: no live EAP, Aurora, or application connection supplied",
              "git_diff_check": "not applicable: workspace is not a Git repository; UTF-8 and whitespace checks passed"}
    (OUT / "validation_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
